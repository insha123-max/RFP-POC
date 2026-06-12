"""Stage 6 — Generate a Word (.docx) evaluation report for committee review."""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT

from models import EvaluationReport


# ── Colour palette ────────────────────────────────────────────────────────────
GREEN  = RGBColor(0x05, 0x96, 0x69)
RED    = RGBColor(0xDC, 0x26, 0x26)
AMBER  = RGBColor(0xD9, 0x77, 0x06)
BLUE   = RGBColor(0x1E, 0x40, 0xAF)
DARK   = RGBColor(0x0F, 0x17, 0x2A)
GREY   = RGBColor(0x64, 0x74, 0x8B)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGREEN = RGBColor(0xDC, 0xFC, 0xE7)
LRED   = RGBColor(0xFE, 0xE2, 0xE2)
LBLUE  = RGBColor(0xEF, 0xF6, 0xFF)


def _shade_cell(cell, hex_color: str):
    """Fill a table cell with a background colour."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def _set_col_widths(table, widths: list[float]):
    """Set column widths in inches."""
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Inches(widths[i])


def _para(doc, text: str, bold=False, size=11, color=None,
          align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold      = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def _heading(doc, text: str, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 11)
    run.font.color.rgb = DARK
    return p


def _table_header_row(table, headers: list[str], bg="1E40AF"):
    row = table.rows[0]
    for i, h in enumerate(headers):
        if i >= len(row.cells):
            break
        cell = row.cells[i]
        _shade_cell(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = WHITE


def generate_word_report(report: EvaluationReport) -> bytes:
    doc  = Document()
    sec  = doc.sections[0]
    sec.page_width   = Inches(8.5)
    sec.page_height  = Inches(11)
    sec.left_margin  = Inches(1)
    sec.right_margin = Inches(1)
    sec.top_margin   = Inches(0.9)
    sec.bottom_margin = Inches(0.9)

    # ── Title block ──────────────────────────────────────────────────────────
    _para(doc, "RFP / TENDER EVALUATION REPORT",
          bold=True, size=18, color=DARK,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    _para(doc, f"Generated on {datetime.now().strftime('%d %B %Y, %H:%M')}",
          size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ── Verdict block ────────────────────────────────────────────────────────
    verdict_text = "✓  PASSED" if report.passed else "✗  FAILED"
    verdict_color = GREEN if report.passed else RED
    _para(doc, verdict_text, bold=True, size=22, color=verdict_color,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    _para(doc,
          f"Total Score: {report.total_score} / {report.max_score}    "
          f"  Threshold: {report.threshold} / {report.max_score}",
          size=11, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    if report.disqualified and report.disqualification_reason:
        _para(doc,
              f"⚠  Automatic Disqualification: {report.disqualification_reason}",
              bold=True, size=10, color=RED,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    doc.add_paragraph()

    # ── Executive Summary ────────────────────────────────────────────────────
    _heading(doc, "Executive Summary")
    _para(doc, report.executive_summary, size=10, color=DARK, space_after=10)

    # ── Category Score Breakdown ─────────────────────────────────────────────
    _heading(doc, "Category Score Breakdown")

    cat_cols = ["Category", "Max Marks", "Awarded", "% Achieved", "Min Required", "Status"]
    t = doc.add_table(rows=1 + len(report.category_results) + 1,
                      cols=len(cat_cols))
    t.style = "Table Grid"
    _table_header_row(t, cat_cols)
    _set_col_widths(t, [2.2, 0.85, 0.85, 0.9, 1.0, 0.8])

    for i, cr in enumerate(report.category_results, start=1):
        row   = t.rows[i]
        cells = row.cells
        data  = [
            cr.category,
            str(cr.max_marks),
            str(cr.marks_awarded),
            f"{cr.percent_achieved}%",
            f"{cr.minimum_required}%" if cr.minimum_required else "—",
            "PASS" if cr.passed else "FAIL",
        ]
        bg = "F0FDF4" if cr.passed else "FEF2F2"
        for j, val in enumerate(data):
            _shade_cell(cells[j], bg)
            p   = cells[j].paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9)
            run.bold = (j == 0 or j == 5)
            if j == 5:
                run.font.color.rgb = GREEN if cr.passed else RED

    # Totals row
    tot  = t.rows[len(report.category_results) + 1]
    totd = ["TOTAL (Weighted)", str(report.max_score), str(report.total_score),
            f"{round(report.total_score/report.max_score*100,1)}%" if report.max_score else "—",
            f"≥ {report.threshold}", "PASSED" if report.passed else "FAILED"]
    for j, val in enumerate(totd):
        _shade_cell(tot.cells[j], "DBEAFE")
        p   = tot.cells[j].paragraphs[0]
        run = p.add_run(val)
        run.bold = True
        run.font.size = Pt(9)
        if j == 5:
            run.font.color.rgb = GREEN if report.passed else RED

    doc.add_paragraph()

    # ── Mandatory Eligibility Checks ─────────────────────────────────────────
    if report.disqualifier_checks:
        _heading(doc, "Mandatory Eligibility Checks")
        dc_cols = ["Condition", "Status", "Notes"]
        dt = doc.add_table(rows=1 + len(report.disqualifier_checks), cols=3)
        dt.style = "Table Grid"
        _table_header_row(dt, dc_cols)
        _set_col_widths(dt, [2.5, 0.8, 3.4])

        for i, chk in enumerate(report.disqualifier_checks, start=1):
            row = dt.rows[i]
            bg  = "F0FDF4" if chk.met else "FEF2F2"
            for j, val in enumerate([chk.condition, "✓ Met" if chk.met else "✗ Not Met", chk.note]):
                _shade_cell(row.cells[j], bg)
                p   = row.cells[j].paragraphs[0]
                run = p.add_run(val)
                run.font.size = Pt(9)
                if j == 1:
                    run.font.color.rgb = GREEN if chk.met else RED
                    run.bold = True

        doc.add_paragraph()

    # ── Criterion-Level Detail ────────────────────────────────────────────────
    _heading(doc, "Criterion-Level Evaluation Details")

    for cr in report.category_results:
        _para(doc, cr.category, bold=True, size=10, color=BLUE, space_before=8, space_after=2)
        if not cr.criteria:
            continue

        crit_cols = ["Criterion", "Max", "Awarded", "Status", "Logic", "Evidence", "Source"]
        ct = doc.add_table(rows=1 + len(cr.criteria), cols=len(crit_cols))
        ct.style = "Table Grid"
        _table_header_row(ct, crit_cols, bg="374151")
        _set_col_widths(ct, [1.6, 0.45, 0.55, 0.7, 0.85, 1.9, 0.9])

        for i, c in enumerate(cr.criteria, start=1):
            row = ct.rows[i]
            bg = "F0FDF4" if c.compliance_status == "Met" else "FEF2F2"
            logic_label = c.threshold_logic or "50% Fallback"
            vals = [c.criterion, str(c.max_marks), str(c.marks_awarded),
                    c.compliance_status, logic_label, c.vendor_claim[:200], c.source_reference]
            for j, val in enumerate(vals):
                _shade_cell(row.cells[j], bg)
                p   = row.cells[j].paragraphs[0]
                run = p.add_run(val)
                run.font.size = Pt(8)
                if j == 3:
                    run.font.color.rgb = GREEN if c.compliance_status == "Met" else RED
                    run.bold = True
                if j == 4:
                    run.font.color.rgb = BLUE if logic_label.startswith("RFP") else GREY
                    run.bold = False

        doc.add_paragraph()

    # ── Gap Analysis & Risk Flags ─────────────────────────────────────────────
    if report.risk_items:
        _heading(doc, "Gap Analysis & Risk Flags")
        risk_cols = ["Risk Area", "Severity", "Description"]
        rt = doc.add_table(rows=1 + len(report.risk_items), cols=3)
        rt.style = "Table Grid"
        _table_header_row(rt, risk_cols, bg="92400E")
        _set_col_widths(rt, [1.6, 0.8, 4.3])

        sev_color = {"High": "FEF2F2", "Medium": "FFFBEB", "Low": "F0FDF4"}
        sev_font  = {"High": RED,       "Medium": AMBER,    "Low": GREEN}

        for i, r in enumerate(report.risk_items, start=1):
            row = rt.rows[i]
            bg  = sev_color.get(r.severity, "FFFFFF")
            for j, val in enumerate([r.risk_area, r.severity, r.description]):
                _shade_cell(row.cells[j], bg)
                p   = row.cells[j].paragraphs[0]
                run = p.add_run(val)
                run.font.size = Pt(9)
                if j == 1:
                    run.font.color.rgb = sev_font.get(r.severity, DARK)
                    run.bold = True

        doc.add_paragraph()

    # ── Footer note ───────────────────────────────────────────────────────────
    _para(doc,
          "This report was auto-generated by the RFP Evaluator AI pipeline. "
          "All scores are based on evidence found in the uploaded documents. "
          "Human review and override are recommended before final award decision.",
          size=8, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER,
          space_before=16, space_after=0)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
